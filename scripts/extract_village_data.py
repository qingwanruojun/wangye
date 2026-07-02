# -*- coding: utf-8 -*-
"""Read 村民培训表 / 村民账号信息 xlsx → JSON / Markdown / Word + 派生指标。"""
from __future__ import annotations

import json
import pathlib
import re
import sys
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data"


def find_xlsx_files() -> tuple[pathlib.Path | None, pathlib.Path | None]:
    training = account = None
    for p in ROOT.glob("*.xlsx"):
        name = p.name
        if "培训" in name:
            training = p
        elif "账号" in name or "信息" in name:
            account = p
    files = sorted(ROOT.glob("*.xlsx"), key=lambda x: x.stat().st_size)
    if len(files) >= 2:
        small, large = files[0], files[-1]
        if small.stat().st_size < 50_000:
            account = account or small
        if large.stat().st_size > 100_000:
            training = training or large
    return training, account


def sheet_to_rows(ws, max_rows: int = 5000) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows.append(list(row))
        if i + 1 >= max_rows:
            break
    return rows


def excel_date_cell(v: Any) -> str | Any:
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        try:
            dt = from_excel(float(v))
            return dt.strftime("%Y-%m-%d") if dt else v
        except Exception:
            return v
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    return v


def derive_training_stats(rows: list[list[Any]]) -> dict[str, Any]:
    """村民培训表 Sheet1：表头行后为场次；签到列为第 4 列（0-based index 3）。"""
    sessions: list[dict[str, Any]] = []
    person_times = 0
    attend_unique: set[str] = set()

    for row in rows[2:]:
        if not row or row[0] is None:
            continue
        sid = row[0]
        if not isinstance(sid, (int, float)):
            continue
        sid_int = int(sid)
        if sid_int < 1:
            continue

        raw_date = excel_date_cell(row[1]) if len(row) > 1 else None
        form = row[2] if len(row) > 2 else None
        sign_raw = row[3] if len(row) > 3 else None
        topic = row[4] if len(row) > 4 else None
        duration = row[5] if len(row) > 5 else None

        names: list[str] = []
        if isinstance(sign_raw, str) and sign_raw.strip():
            for part in re.split(r"[\n\r,，、]+", sign_raw.strip()):
                n = part.strip()
                if n:
                    names.append(n)
                    attend_unique.add(n)

        person_times += len(names)
        sessions.append(
            {
                "场次序号": sid_int,
                "日期": raw_date,
                "形式": form,
                "签到人数（本场）": len(names),
                "签到名单": names,
                "培训内容": topic if isinstance(topic, str) else topic,
                "课时": duration,
            }
        )

    return {
        "培训场次数": len(sessions),
        "累计培训人次（签到合计）": person_times,
        "签到涉及人数（去重）": len(attend_unique),
        "场次明细": sessions,
    }


def derive_account_stats(rows: list[list[Any]]) -> dict[str, Any]:
    """账号表：第 1 行为总标题，第 2 行列名，第 3 行最近视频子列名，数据从第 4 行起。"""
    accounts: list[dict[str, Any]] = []
    total_original = 0
    max_play = 0

    for row in rows[3:]:
        if not row or row[0] is None:
            continue
        try:
            seq = int(row[0])
        except (TypeError, ValueError):
            continue
        name = row[1] if len(row) > 1 else None
        if name is None or name == "":
            continue

        orig = row[2] if len(row) > 2 else None
        if isinstance(orig, (int, float)):
            total_original += int(orig)

        plays: list[int] = []
        for idx in (6, 9):
            if len(row) > idx and isinstance(row[idx], (int, float)):
                pv = int(row[idx])
                plays.append(pv)
                max_play = max(max_play, pv)

        accounts.append(
            {
                "序号": seq,
                "账号名称": name,
                "原创视频数": orig,
                "关注量": row[3] if len(row) > 3 else None,
                "点赞量": row[4] if len(row) > 4 else None,
                "收藏量": row[5] if len(row) > 5 else None,
                "近期播放量样本": plays,
            }
        )

    return {
        "台账账号数": len(accounts),
        "原创视频数合计（有填报行）": total_original,
        "近期样本最大播放量": max_play if max_play else None,
        "账号明细": accounts,
    }


def trim_rows(rows: list[list[Any]], max_keep: int = 120) -> list[list[Any]]:
    nonempty = [r for r in rows if any(c is not None and c != "" for c in r)]
    if len(nonempty) <= max_keep:
        return [[excel_date_cell(c) for c in r] for r in nonempty]
    head = [[excel_date_cell(c) for c in r] for r in nonempty[: max_keep - 1]]
    head.append(["…", f"(省略 {len(nonempty) - (max_keep - 1)} 行空行或尾部)", "…"])
    return head


def main():
    training_path, account_path = find_xlsx_files()
    if not training_path or not training_path.exists():
        print("ERROR: 未找到村民培训表 xlsx", file=sys.stderr)
        sys.exit(1)
    if not account_path or not account_path.exists():
        print("ERROR: 未找到村民账号信息 xlsx", file=sys.stderr)
        sys.exit(1)

    OUT_DIR.mkdir(exist_ok=True)

    wb_acc = openpyxl.load_workbook(account_path, read_only=True, data_only=True)
    account_payload: dict[str, Any] = {"source_file": account_path.name, "sheets": {}}
    account_sheet1_rows: list[list[Any]] = []
    for sn in wb_acc.sheetnames:
        ws = wb_acc[sn]
        rows = sheet_to_rows(ws, max_rows=500)
        account_payload["sheets"][sn] = {
            "row_count": len(rows),
            "rows": trim_rows(rows, max_keep=80),
        }
        if sn == "Sheet1":
            account_sheet1_rows = rows
    wb_acc.close()

    wb_tr = openpyxl.load_workbook(training_path, read_only=True, data_only=True)
    training_payload: dict[str, Any] = {"source_file": training_path.name, "sheets": {}}
    training_sheet1_rows: list[list[Any]] = []
    for sn in wb_tr.sheetnames:
        ws = wb_tr[sn]
        rows = sheet_to_rows(ws, max_rows=4000)
        training_payload["sheets"][sn] = {
            "row_count_sampled": len(rows),
            "truncated": len(rows) >= 4000,
            "rows": trim_rows(rows, max_keep=100),
        }
        if sn == "Sheet1":
            training_sheet1_rows = rows
    wb_tr.close()

    derived = {
        "培训统计": derive_training_stats(training_sheet1_rows),
        "账号统计": derive_account_stats(account_sheet1_rows),
    }

    bundle = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "derived": derived,
        "training": training_payload,
        "account": account_payload,
    }

    json_path = OUT_DIR / "village_excel_export.json"
    json_path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")

    derived_only = OUT_DIR / "village_derived_stats.json"
    derived_only.write_text(
        json.dumps(
            {
                "exported_at": bundle["exported_at"],
                "sources": {
                    "training": training_path.name,
                    "account": account_path.name,
                },
                **derived,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    md_lines = [
        "# 村民台账导出",
        "",
        f"- 导出时间：`{bundle['exported_at']}`",
        f"- 培训表：`{training_path.name}`",
        f"- 账号信息：`{account_path.name}`",
        "",
        "## 派生指标（供大屏 / 主站对齐）",
        "",
        f"- **培训场次数**：{derived['培训统计']['培训场次数']}",
        f"- **累计培训人次（签到合计）**：{derived['培训统计']['累计培训人次（签到合计）']}",
        f"- **签到涉及人数（去重）**：{derived['培训统计']['签到涉及人数（去重）']}",
        f"- **台账账号数**：{derived['账号统计']['台账账号数']}",
        f"- **原创视频数合计**：{derived['账号统计']['原创视频数合计（有填报行）']}",
        f"- **近期样本最大播放量**：{derived['账号统计']['近期样本最大播放量'] or '—'}",
        "",
        "## 账号信息 · 工作表概览",
        "",
    ]
    for sn, info in account_payload["sheets"].items():
        md_lines.append(f"- **{sn}**：{info['row_count']} 行")
    md_lines.extend(["", "## 培训表 · 工作表概览", ""])
    for sn, info in training_payload["sheets"].items():
        flag = "（已截断抽样）" if info.get("truncated") else ""
        md_lines.append(f"- **{sn}**：{info['row_count_sampled']} 行{flag}")

    md_path = OUT_DIR / "village_excel_summary.md"
    md_path.write_text("\n".join(md_lines), encoding="utf-8")

    # --- Word ---
    docx_path = OUT_DIR / "village_excel_summary.docx"
    from docx import Document  # type: ignore

    doc = Document()
    doc.add_heading("村民台账导出（培训表 + 账号信息）", level=0)
    doc.add_paragraph(f"导出时间：{bundle['exported_at']}")
    doc.add_paragraph(f"源文件：{training_path.name}；{account_path.name}")

    doc.add_heading("一、派生指标", level=1)
    for label, keypath in (
        ("培训场次数", ("培训统计", "培训场次数")),
        ("累计培训人次（签到合计）", ("培训统计", "累计培训人次（签到合计）")),
        ("签到涉及人数（去重）", ("培训统计", "签到涉及人数（去重）")),
        ("台账账号数", ("账号统计", "台账账号数")),
        ("原创视频数合计（有填报行）", ("账号统计", "原创视频数合计（有填报行）")),
        ("近期样本最大播放量", ("账号统计", "近期样本最大播放量")),
    ):
        ps = derived
        for k in keypath:
            ps = ps[k]
        doc.add_paragraph(f"{label}：{ps if ps is not None else '—'}", style="List Bullet")

    doc.add_heading("二、培训场次明细", level=1)
    t = doc.add_table(rows=1, cols=5)
    hdr = t.rows[0].cells
    hdr[0].text = "序号"
    hdr[1].text = "日期"
    hdr[2].text = "形式"
    hdr[3].text = "本场人次"
    hdr[4].text = "内容摘要"
    for s in derived["培训统计"]["场次明细"]:
        row = t.add_row().cells
        row[0].text = str(s["场次序号"])
        row[1].text = str(s["日期"] or "")
        row[2].text = str(s["形式"] or "")
        row[3].text = str(s["签到人数（本场）"])
        topic = s.get("培训内容") or ""
        row[4].text = (
            topic.replace("\n", " ")[:120] + ("…" if isinstance(topic, str) and len(topic) > 120 else "")
            if isinstance(topic, str)
            else str(topic)
        )

    doc.add_heading("三、账号台账摘录", level=1)
    t2 = doc.add_table(rows=1, cols=5)
    h2 = t2.rows[0].cells
    h2[0].text = "序号"
    h2[1].text = "账号名称"
    h2[2].text = "原创数"
    h2[3].text = "点赞"
    h2[4].text = "近期播放样本"
    for a in derived["账号统计"]["账号明细"]:
        row = t2.add_row().cells
        row[0].text = str(a["序号"])
        row[1].text = str(a["账号名称"])
        row[2].text = "" if a["原创视频数"] is None else str(a["原创视频数"])
        row[3].text = "" if a["点赞量"] is None else str(a["点赞量"])
        row[4].text = "/".join(str(x) for x in a["近期播放量样本"]) if a["近期播放量样本"] else "—"

    doc.save(docx_path)

    print("OK")
    print(" JSON:", json_path)
    print(" JSON:", derived_only)
    print(" MD :", md_path)
    print(" DOCX:", docx_path)


if __name__ == "__main__":
    main()
